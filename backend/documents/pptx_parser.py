from pptx import Presentation


def extract_text(file_path: str) -> str:
    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        if slide_text:
            slides.append(
                f"Slide {slide_number}\n"
                + "\n".join(slide_text)
            )

    return "\n\n".join(slides).strip()